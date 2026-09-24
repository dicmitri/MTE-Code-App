#!/usr/bin/env python3
"""Build the editable MedTech Europe Code DOCX directly from the source PDF.

The PDF is the only content input. The split JSON chapters are deliberately not
read or imported. Approved source typographical errors are corrected through the
explicit, page-scoped CORRECTIONS list below and recorded in the DOCX appendix.

Each page is read per column with pdfplumber, so paragraphs, list levels, bold
words and superscript footnote references follow the PDF's layout. Letter, roman
and number labels are typed text, exactly as printed; only bullets use Word lists.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

import pdfplumber
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


PROJECT_ROOT = Path(__file__).resolve().parent.parent
SOURCE_PDF = PROJECT_ROOT / "src" / "data" / "code-september-2024 (1).pdf"
OUTPUT_DOCX = PROJECT_ROOT / "src" / "data" / "code-september-2024.docx"
SOURCE_SHA256 = "9ED658D0C9F858E35576C3905CC4423DF934DB894786F4C12D9A130E0602A148"

INK = RGBColor(0x20, 0x2A, 0x35)
MUTED = RGBColor(0x5E, 0x68, 0x75)
HEADING_BLUE = RGBColor(0x2E, 0x74, 0xB5)
HEADING_DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
QUESTION_PURPLE = RGBColor(0x76, 0x54, 0xA1)
TEAL = RGBColor(0x00, 0x7A, 0x86)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
NOTE_FILL = "F4F6F9"
TABLE_BORDER = "AAB4C0"


@dataclass(frozen=True)
class Correction:
    page: int
    pattern: str
    replacement: str
    original: str
    reason: str
    flags: int = 0


CORRECTIONS = (
    Correction(7, r"\bsponsorhip\b", "sponsorship", "sponsorhip", "Spelling"),
    Correction(
        7,
        r"\bProfessionalsthe\b",
        "Professionals the",
        "Professionalsthe",
        "Joined words",
    ),
    Correction(
        8,
        r"from the date of\s+the date of the change",
        "from the date of the change",
        "from the date of the date of the change",
        "Duplicated words",
    ),
    Correction(
        8,
        r"Assembly\.\.",
        "Assembly.",
        "Assembly..",
        "Duplicated punctuation",
    ),
    Correction(
        16,
        r"\bMedicalTechnologies\b",
        "Medical Technologies",
        "MedicalTechnologies",
        "Joined words",
    ),
    Correction(
        18,
        r"\bwith the the requirements\b",
        "with the requirements",
        "with the the requirements",
        "Duplicated word",
    ),
    Correction(
        22,
        r"\bCompanyMedical\b",
        "Company Medical",
        "CompanyMedical",
        "Joined words",
    ),
    Correction(
        26,
        r"where the meeting is being hosted$",
        "where the meeting is being hosted.",
        "where the meeting is being hosted",
        "Missing sentence-ending punctuation",
    ),
    Correction(
        26,
        r"a larger Third Party Organised Educational Conferences",
        "a larger Third Party Organised Educational Conference",
        "a larger Third Party Organised Educational Conferences",
        "Singular agreement",
    ),
    Correction(27, r"\bKkind\b", "Kind", "Kkind", "Spelling"),
    Correction(35, r"\bTechnoloy\b", "Technology", "Technoloy", "Spelling"),
    Correction(
        33,
        r"Charitable Donations\.\.",
        "Charitable Donations.",
        "Charitable Donations..",
        "Duplicated punctuation",
    ),
    Correction(
        37,
        r"However,Member Ccompanies",
        "However, Member Companies",
        "However,Member Ccompanies",
        "Spacing and spelling",
    ),
    Correction(37, r"\bcritera\b", "criteria", "critera", "Spelling"),
    Correction(
        39,
        r"\bCompanyMedical\b",
        "Company Medical",
        "CompanyMedical",
        "Joined words",
    ),
    Correction(
        40,
        r"\bCcompany-organised\b",
        "Company-organised",
        "Ccompany-organised",
        "Spelling",
    ),
    Correction(
        38,
        r"Healthcare Organisation, , all",
        "Healthcare Organisation, all",
        "Healthcare Organisation, , all",
        "Duplicated punctuation",
    ),
    Correction(
        43,
        r"Medical Technology,\.",
        "Medical Technology.",
        "Medical Technology,.",
        "Duplicated punctuation",
    ),
    Correction(
        44,
        r"scope of the Consultancy Arrangement\s*Member Companies",
        "scope of the Consultancy Arrangement. Member Companies",
        "scope of the Consultancy Arrangement Member Companies",
        "Missing sentence-ending punctuation",
    ),
    Correction(
        47,
        r"\bwww\.who\.org\b",
        "www.who.int.",
        "www.who.org",
        "Corrected URL and sentence-ending punctuation",
    ),
    Correction(
        48,
        r"\bethicalrequirements\b",
        "ethical requirements",
        "ethicalrequirements",
        "Joined words",
    ),
    Correction(49, r"\basFaculty\b", "as Faculty", "asFaculty", "Joined words"),
    Correction(
        49,
        r"\(“Glossary’\)",
        "(“Glossary”)",
        "(“Glossary’)",
        "Mismatched quotation mark",
    ),
    Correction(49, r"\byhird party-", "third party-", "yhird party-", "Spelling"),
    Correction(
        49,
        r"\bMember companies may contract researchers\b",
        "Member Companies may contract researchers",
        "Member companies may contract researchers",
        "Defined-term capitalization",
    ),
    Correction(
        50,
        r"\bMember\s+company\s+to\s+conduct\s+collaborative\s+research\b",
        "Member Company to conduct collaborative research",
        "Member company to conduct collaborative research",
        "Defined-term capitalization",
    ),
    Correction(
        50,
        r"\bis clear expressed\b",
        "is clearly expressed",
        "is clear expressed",
        "Grammar",
    ),
    Correction(
        54,
        r"\bChapter 8\.Educational\b",
        "Chapter 8. Educational",
        "Chapter 8.Educational",
        "Missing space",
    ),
    Correction(54, r"\bPromotionl\b", "Promotional", "Promotionl", "Spelling"),
    Correction(57, r"\bMedica;l\b", "Medical", "Medica;l", "Spelling"),
    Correction(
        61,
        r"\bMember companies should encourage\b",
        "Member Companies should encourage",
        "Member companies should encourage",
        "Defined-term capitalization",
    ),
    Correction(
        61,
        r"\bapplicable applicable\b",
        "applicable",
        "applicable applicable",
        "Duplicated word",
    ),
    Correction(
        64,
        r"\blaid down in in this section\b",
        "laid down in this section",
        "laid down in in this section",
        "Duplicated word",
    ),
    Correction(
        71,
        r"\ba Members Company’s medical\b",
        "a Member Company’s medical",
        "a Members Company’s medical",
        "Defined-term spelling",
    ),
    Correction(
        73,
        r"\bdirectlyby\b",
        "directly by",
        "directlyby",
        "Joined words",
    ),
    Correction(
        72,
        r"within this category$",
        "within this category.",
        "within this category",
        "Missing sentence-ending punctuation",
    ),
    Correction(
        76,
        r"agreement, the Member Company\.$",
        "agreement, the Member Company:",
        "agreement, the Member Company.",
        "List-introducing punctuation",
    ),
    Correction(
        77,
        r"\bcountries specificities\b",
        "country specificities",
        "countries specificities",
        "Grammar",
    ),
    Correction(
        78,
        r"\bThird\s+Third\s+Party Organised Educational Event\b",
        "Third Party Organised Educational Event",
        "Third Third Party Organised Educational Event",
        "Duplicated word",
    ),
    Correction(
        78,
        r"\bit is a thirdpParty chooses\b",
        "a third party chooses",
        "it is a thirdpParty chooses",
        "Joined word and grammar",
    ),
    Correction(
        78,
        r"\bthird pParty\b",
        "third party",
        "third pParty",
        "Spelling",
    ),
    Correction(79, r"\b3D rending\b", "3D rendering", "3D rending", "Spelling"),
)


MAJOR_TITLES = {
    5: "Scope",
    9: "Administering the Code",
    14: "Introduction",
    18: "Chapter 1: General Criteria for Event",
    24: "Chapter 2: Third Party Organised Educational Events",
    28: "Chapter 3: Company Events",
    33: "Chapter 4: Grants and Charitable Donations",
    41: "Chapter 5: Consulting Arrangements",
    46: "Chapter 6: Research",
    51: "Chapter 7: Royalties",
    53: "Chapter 8: Educational Items and Promotional Items",
    56: "Chapter 9: Demonstration Products and Samples",
    59: "Chapter 10: Third Party Intermediaries",
    62: (
        "MedTech Europe Code of Ethical Business Practice Part 2: "
        "Complaint handling and dispute resolution"
    ),
    67: (
        "MedTech Europe Code of Ethical Business Practice Part 3: "
        "Glossary and Definitions"
    ),
}

COVER_BODY_STARTS = {
    9: "Part 2 of the Code includes",
    18: "Member Companies may directly finance",
    24: "Member Companies may provide financial",
}

ADDITIONAL_HEADINGS = {
    "About MedTech Europe",
    "Promoting a balanced policy environment",
    "Demonstrating the value of medical technology",
    "Promoting an Ethical Industry",
    "Key Legislation",
    "Aims and Principles of the Code",
    "1. Tasks",
    "2. Composition",
    "3. MedTech Europe Compliance Panel Internal Procedural Rules",
    "2.1 New Member Companies",
    "2.2 New Association Members",
    "Advancement of Medical Technologies",
    "Safe and Effective Use of Medical Technology",
    "Research and Education",
    "Risk Assessment",
    "Due diligence",
    "Training",
    "Written Contract",
    "Oversight",
    "Appropriate Corrective Action",
    "Description:",
    "Guidance:",
    "What is an In-Kind Educational Grant?",
    "Types of In-Kind Educational Grant",
    "Value of In Kind Contributions",
    "Countries with National Associations:",
    (
        "Countries party to the European Economic Area agreement without a "
        "MedTech Europe National Association:"
    ),
    "Verification Of The Use Of Funds",
    "Annex IV",
    "Annex V",
    "ANNEX VII",
    (
        "The Criteria Applicable to Third Party Organised Procedure Trainings "
        "(Effective as of 3rd September 2018)"
    ),
    (
        "How can a Member Company verify that the Educational Grant is in fact "
        "used for the intended purpose as agreed in the Educational Grant agreement?"
    ),
    (
        "Grant to support Healthcare Professionals’ attendance at the Third Party "
        "Organised Educational Event:"
    ),
    (
        "Grant to support the costs related to organisation of the Third Party "
        "Organised Educational Event:"
    ),
    "Grant provided in a form of a Scholarship or Fellowship:",
    "Methodology Note Example",
    "Structure",
    "Disclosure scope and timelines",
    "Disclosures in case of partial performance or cancellation",
    "Cross-border activities",
    "Specific considerations:",
    "The Criteria Applicable to Third Party Organised Procedure Trainings",
    "CRITERIA FOR TPPT DETERMINATION",
    "1. Programme:",
    "2. Venue:",
    "3. Stand-alone event:",
    "4. Size:",
}

TERMINAL_RE = re.compile(r'[.!?;:](?:["”’\)])?$')
QUESTION_RE = re.compile(r"^Q\d+\b")
ANSWER_RE = re.compile(r"^A\d+\b")
CLAUSE_RE = re.compile(r"^\d+\.\d+(?:\.\d+)?\.?\s")
FOOTNOTE_RE = re.compile(r"^\d+\)\s+")
# Q&A labels the PDF prints beside a wrapped question or answer, between its lines.
MARGIN_LABEL_RE = re.compile(r"[QA]\d+")
# Letter, roman and number labels are kept as typed text, exactly as printed.
LABEL_RE = re.compile(
    r"^(\(?[a-z][.)]|\(?(?:i{1,3}|iv|vi{0,3}|ix|x)[.)]?|\d{1,2}[.)])\s+(?=\S)"
)
BULLET_GLYPHS = {"•": "bullet", "■": "square", "-": "dash"}
BULLET_KINDS = frozenset(BULLET_GLYPHS.values())
LIST_KINDS = BULLET_KINDS | {"label"}
TEXT_KINDS = LIST_KINDS | {"body", "clause", "question", "answer", "footnote"}

# PDF layout (points). Page furniture sits above BODY_TOP and below BODY_BOTTOM;
# two-column pages print the Q&As to the right of SIDEBAR_X.
BODY_TOP = 30.0
BODY_BOTTOM = 565.0
SIDEBAR_X = 455.0
PARAGRAPH_GAP = 3.2  # extra vertical space that starts a new paragraph
SHORT_LINE = 12.0  # a line ending this far before the column edge can end a paragraph
FULL_LINE = 3.0  # a line reaching the column edge continues on the next line
LEVEL_STEP = 4.0  # indentation difference that marks a nested list level
FOOTNOTE_SIZE = 7.25  # footnotes print at 6-7pt, Q&As at 7.5pt and body text at 9pt

# Annex V prints a disclaimer box beside its last list items.
REGION_OVERRIDES = {
    77: (
        ("main", (0, BODY_TOP, None, 495)),
        ("main", (0, 495, 440, BODY_BOTTOM)),
        ("note", (440, 495, None, BODY_BOTTOM)),
    ),
}

# Inline markup carried from the PDF glyphs to the Word runs.
SUP_START, SUP_END = "\x01", "\x02"
BOLD_START, BOLD_END = "\x03", "\x04"
MARKUP_RE = re.compile("[\x01-\x04]")


@dataclass
class Line:
    text: str
    x0: float
    x1: float
    top: float
    bottom: float
    right: float
    font: str
    size: float


@dataclass
class Block:
    kind: str
    text: str
    page: int
    level: int = 0
    column: str = "main"
    x0: float = 0.0
    label: str = ""
    ends_full: bool = False
    indent: int | None = None
    pages: list[int] = field(default_factory=list)
    bold: bool = False
    line_count: int = 1


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_space(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text).strip()


def plain(text: str) -> str:
    return MARKUP_RE.sub("", text)


def join_wrapped(left: str, right: str) -> str:
    left = left.rstrip()
    right = right.lstrip()
    if not left:
        return right
    if not right:
        return left
    if plain(left).endswith(("-", "/", "–")):
        joined = left + right
    else:
        joined = f"{left} {right}"
    # A bold span that wraps onto the next line stays one span.
    return re.sub(f"{BOLD_END}(\\s*){BOLD_START}", r"\1", joined)


def strip_visible(text: str, count: int) -> str:
    """Drop the first `count` visible characters, reopening any markup they open."""
    kept: list[str] = []
    opened: list[str] = []
    seen = 0
    for character in text:
        if MARKUP_RE.match(character):
            if seen < count:
                if character in (SUP_START, BOLD_START):
                    opened.append(character)
                elif opened:
                    opened.pop()
            else:
                kept.append(character)
            continue
        if seen < count:
            seen += 1
            continue
        kept.append(character)
    return "".join(opened) + "".join(kept).lstrip()


class CorrectionLog:
    """Applies CORRECTIONS and proves each one matched exactly once in the build."""

    def __init__(self) -> None:
        self.counts = {correction: 0 for correction in CORRECTIONS}

    def apply(self, page_number: int, text: str) -> str:
        for correction in CORRECTIONS:
            if correction.page != page_number:
                continue
            text, count = re.subn(
                correction.pattern,
                correction.replacement,
                text,
                flags=correction.flags,
            )
            self.counts[correction] += count
        return text

    def assert_complete(self) -> None:
        wrong = [
            f"page {correction.page}: {correction.original!r} matched {count} times"
            for correction, count in self.counts.items()
            if count != 1
        ]
        if wrong:
            raise ValueError("Corrections must match exactly once: " + "; ".join(wrong))


def extract_toc_headings(plumber_document) -> set[str]:
    """Headings named in the PDF's own table of contents (pages 3-4)."""
    headings = set(ADDITIONAL_HEADINGS)
    for page_number in (3, 4):
        page = plumber_document.pages[page_number - 1]
        body = page.crop((0, BODY_TOP, float(page.width), BODY_BOTTOM))
        for line in (body.extract_text() or "").splitlines():
            match = re.match(r"^(.*?)\s+\d+$", normalize_space(line))
            if match:
                headings.add(match.group(1))
    return headings


def styled_text(line: dict) -> tuple[str, str, float]:
    """Line text with superscript and bold markup, the line's font kind and size."""
    glyphs = [char for char in line["chars"] if char["text"].strip()]
    text = line["text"]
    if not glyphs:
        return text, "normal", 0.0
    body_size = statistics.median(char["size"] for char in glyphs)
    marked: list[str] = []
    index = 0
    superscript = bold = False
    for character in text:
        if character.isspace():
            if superscript:
                marked.append(SUP_END)
                superscript = False
            marked.append(character)
            continue
        glyph = glyphs[index]
        index += 1
        if glyph["text"] != character:
            raise ValueError(f"Glyph order differs from line text: {text!r}")
        is_superscript = glyph["size"] < body_size * 0.75 and (
            character.isdigit() or character == ","
        )
        is_bold = "Bold" in glyph["fontname"] and not is_superscript
        if is_bold != bold:
            marked.append(BOLD_START if is_bold else BOLD_END)
            bold = is_bold
        if is_superscript != superscript:
            marked.append(SUP_START if is_superscript else SUP_END)
            superscript = is_superscript
        marked.append(character)
    if superscript:
        marked.append(SUP_END)
    if bold:
        marked.append(BOLD_END)
    # Spaces after a bold span belong to the following text.
    result = re.sub(f"(\\s+){BOLD_END}", f"{BOLD_END}\\1", "".join(marked))
    fonts = {glyph["fontname"].split("+")[-1] for glyph in glyphs}
    if all(font.startswith("Chalet") for font in fonts) and body_size >= 12:
        return result, "title", body_size
    if all("Bold" in font for font in fonts):
        return result, "bold", body_size
    return result, "normal", body_size


def page_regions(page, page_number: int) -> list[tuple[str, tuple[float, float, float, float]]]:
    width = float(page.width)
    if page_number in REGION_OVERRIDES:
        return [
            (column, (x0, top, width if x1 is None else x1, bottom))
            for column, (x0, top, x1, bottom) in REGION_OVERRIDES[page_number]
        ]
    words = page.crop((0, BODY_TOP, width, BODY_BOTTOM)).extract_words()
    two_column = any(word["x0"] >= SIDEBAR_X for word in words) and not any(
        word["x0"] < SIDEBAR_X < word["x1"] for word in words
    )
    if two_column:
        return [
            ("main", (0, BODY_TOP, SIDEBAR_X, BODY_BOTTOM)),
            ("side", (SIDEBAR_X, BODY_TOP, width, BODY_BOTTOM)),
        ]
    return [("main", (0, BODY_TOP, width, BODY_BOTTOM))]


def page_columns(page, page_number: int) -> dict[str, list[Line]]:
    columns: dict[str, list[Line]] = {}
    for column, bbox in page_regions(page, page_number):
        raw = [
            line
            for line in page.crop(bbox).extract_text_lines(strip=True, return_chars=True)
            if line["text"].strip()
        ]
        right = max((line["x1"] for line in raw), default=bbox[2])
        lines = columns.setdefault(column, [])
        for line in raw:
            text, font, size = styled_text(line)
            lines.append(
                Line(
                    text, line["x0"], line["x1"], line["top"], line["bottom"], right, font, size
                )
            )
    return {column: attach_margin_labels(lines) for column, lines in columns.items()}


def attach_margin_labels(lines: Sequence[Line]) -> list[Line]:
    """Put a Q&A label printed beside its wrapped question or answer at its first line."""
    labels = [line for line in lines if MARGIN_LABEL_RE.fullmatch(plain(line.text))]
    if not labels:
        return list(lines)
    label_ids = {id(label) for label in labels}
    text_lines = [line for line in lines if id(line) not in label_ids]
    groups = paragraph_groups(text_lines)
    for label in labels:
        centre = (label.top + label.bottom) / 2
        group = min(
            groups,
            key=lambda group: max(group[0].top - centre, centre - group[-1].bottom, 0),
        )
        first = group[0]
        if line_start(first) is not None:
            raise ValueError(f"Q&A label {label.text!r} has no unlabelled paragraph")
        first.text = f"{plain(label.text)} {first.text}"
        first.x0 = label.x0
    return text_lines


def strip_cover_lines(page_number: int, lines: Sequence[Line]) -> list[Line]:
    body_start = COVER_BODY_STARTS.get(page_number)
    if body_start is None:
        return []
    for index, line in enumerate(lines):
        if plain(line.text).startswith(body_start):
            # Title pages also carry a small hidden copy of the title among their lines.
            return [
                line
                for line in lines[index:]
                if normalize_space(plain(line.text)) != MAJOR_TITLES[page_number]
            ]
    raise ValueError(f"Page {page_number}: cover body start not found: {body_start}")


def line_start(line: Line) -> tuple[str, str, str] | None:
    """(kind, label, remaining text) when a line opens a new item, else None."""
    text = plain(line.text)
    if QUESTION_RE.match(text):
        return "question", "", line.text
    if ANSWER_RE.match(text):
        return "answer", "", line.text
    if text[:1] in BULLET_GLYPHS and (len(text) == 1 or text[1].isspace()):
        return BULLET_GLYPHS[text[0]], "", strip_visible(line.text, min(len(text), 2))
    if FOOTNOTE_RE.match(text) and line.size < FOOTNOTE_SIZE:
        return "footnote", "", line.text
    if CLAUSE_RE.match(text):
        return "clause", "", line.text
    match = LABEL_RE.match(text)
    if match:
        return "label", match.group(1), strip_visible(line.text, match.end())
    return None


def ends_paragraph(line: Line) -> bool:
    return line.right - line.x1 > SHORT_LINE and bool(
        TERMINAL_RE.search(plain(line.text).rstrip())
    )


def paragraph_groups(lines: Sequence[Line]) -> list[list[Line]]:
    groups: list[list[Line]] = []
    previous: Line | None = None
    for line in lines:
        new_group = (
            previous is None
            or line_start(line) is not None
            or line.top - previous.bottom > PARAGRAPH_GAP
            or ends_paragraph(previous)
            or (previous.font == "title") != (line.font == "title")
            # A short bold line is a title above its paragraph.
            or (
                previous.font == "bold"
                and line.font != "bold"
                and previous.right - previous.x1 > SHORT_LINE
            )
        )
        if new_group:
            groups.append([line])
        else:
            groups[-1].append(line)
        previous = line
    return groups


def heading_key(text: str) -> str:
    return normalize_space(plain(text)).casefold()


def parse_column(
    page_number: int,
    column: str,
    lines: Sequence[Line],
    headings: set[str],
) -> list[Block]:
    known_headings = {heading_key(heading) for heading in headings}
    blocks: list[Block] = []
    for group in paragraph_groups(lines):
        first = group[0]
        joined = ""
        for line in group:
            joined = join_wrapped(joined, line.text)
        ends_full = group[-1].right - group[-1].x1 < FULL_LINE
        is_title = all(line.font == "title" for line in group)
        if is_title or heading_key(joined) in known_headings:
            blocks.append(
                Block(
                    "heading",
                    normalize_space(plain(joined)),
                    page_number,
                    level=2 if is_title else 3,
                    column=column,
                    x0=first.x0,
                    ends_full=ends_full,
                )
            )
            continue
        start = line_start(first)
        if start is None:
            kind, label, text = "body", "", joined
        else:
            kind, label, first_text = start
            text = first_text
            for line in group[1:]:
                text = join_wrapped(text, line.text)
        blocks.append(
            Block(
                kind,
                text,
                page_number,
                column=column,
                x0=first.x0,
                label=label,
                ends_full=ends_full,
                bold=all(line.font == "bold" for line in group),
                line_count=len(group),
            )
        )
    return blocks


def promote_subheadings(blocks: Iterable[Block]) -> None:
    """A bold line or two without closing punctuation is a sub-heading."""
    for block in blocks:
        if (
            block.kind in {"body", "clause"}
            and block.bold
            and block.line_count <= 2
            and not re.search(r"[.,;:]$", plain(block.text).rstrip())
        ):
            block.kind = "heading"
            block.level = 3
            block.text = normalize_space(plain(block.text))


def merge_page_breaks(
    pages: dict[int, dict[str, list[Block]]], fresh_starts: set[int]
) -> None:
    """Rejoin paragraphs and headings that the PDF breaks across a page."""
    numbers = sorted(pages)
    for previous_page, page in zip(numbers, numbers[1:]):
        if page != previous_page + 1 or page in fresh_starts:
            continue
        for column in ("main", "side"):
            before = pages[previous_page].get(column)
            after = pages[page].get(column)
            if not before or not after:
                continue
            # Footnotes print below the text that runs on to the next page.
            text_blocks = [block for block in before if block.kind != "footnote"]
            if not text_blocks:
                continue
            last, first = text_blocks[-1], after[0]
            continues_text = (
                last.kind in TEXT_KINDS
                and first.kind == "body"
                and (last.ends_full or not TERMINAL_RE.search(plain(last.text).rstrip()))
            )
            continues_heading = (
                last.kind == first.kind == "heading"
                and last.level == first.level
                and last.ends_full
            )
            if continues_text or continues_heading:
                last.text = join_wrapped(last.text, first.text)
                last.ends_full = first.ends_full
                last.pages.extend([page, *first.pages])
                last.bold = last.bold and first.bold
                last.line_count += first.line_count
                after.pop(0)


def column_text_left(lines: Sequence[Line]) -> float:
    """The x position most lines of a column start at."""
    positions = [round(line.x0) for line in lines if line_start(line) is None]
    if not positions:
        return min((line.x0 for line in lines), default=0.0)
    return float(statistics.mode(positions))


def assign_levels(
    blocks: Sequence[Block], text_left: dict[tuple[int, str], float]
) -> None:
    """Nest list items by their printed indentation within each run of a list."""
    stack: list[float] = []
    for block in blocks:
        if block.kind in LIST_KINDS:
            while stack and block.x0 < stack[-1] - LEVEL_STEP:
                stack.pop()
            if not stack or block.x0 > stack[-1] + LEVEL_STEP:
                stack.append(block.x0)
            block.level = min(len(stack) - 1, 2)
            continue
        left = text_left.get((block.page, block.column))
        if block.kind == "body" and left is not None and block.x0 > left + LEVEL_STEP:
            # Paragraphs printed under a list item keep the list's indentation.
            block.indent = min(max(len(stack) - 1, 0), 2)
            continue
        stack = []


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_style(
    style,
    *,
    size: float,
    color: RGBColor,
    bold: bool = False,
    italic: bool = False,
    before: float = 0,
    after: float = 6,
    line_spacing: float = 1.25,
) -> None:
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:ascii"), "Calibri"
    )
    style._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:hAnsi"), "Calibri"
    )
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style.font.italic = italic
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing


def ensure_paragraph_style(document: Document, name: str):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def configure_styles(document: Document) -> None:
    configure_style(
        document.styles["Normal"], size=11, color=INK, after=6, line_spacing=1.25
    )
    configure_style(
        document.styles["Heading 1"],
        size=16,
        color=HEADING_BLUE,
        bold=True,
        before=18,
        after=10,
        line_spacing=1.0,
    )
    configure_style(
        document.styles["Heading 2"],
        size=13,
        color=HEADING_BLUE,
        bold=True,
        before=14,
        after=7,
        line_spacing=1.0,
    )
    configure_style(
        document.styles["Heading 3"],
        size=12,
        color=HEADING_DARK_BLUE,
        bold=True,
        before=10,
        after=5,
        line_spacing=1.0,
    )
    for heading_name in ("Heading 1", "Heading 2", "Heading 3"):
        document.styles[heading_name].paragraph_format.keep_with_next = True

    styles = {
        "Code Body": dict(size=11, color=INK),
        "Code Clause": dict(size=11, color=INK, after=8),
        "Code Question": dict(
            size=10.5, color=QUESTION_PURPLE, bold=True, before=8, after=3
        ),
        "Code Answer": dict(size=10.5, color=INK, before=0, after=7),
        "Code Footnote": dict(size=9, color=MUTED, after=4, line_spacing=1.1),
        "Code List": dict(size=11, color=INK, after=4),
        "Code Answer List": dict(size=10.5, color=INK, after=4),
        "Editorial Note": dict(
            size=10, color=MUTED, italic=True, after=6, line_spacing=1.15
        ),
        "Table Text": dict(size=9, color=INK, after=2, line_spacing=1.1),
        "Table Header": dict(
            size=9, color=INK, bold=True, after=2, line_spacing=1.05
        ),
    }
    for name, tokens in styles.items():
        configure_style(ensure_paragraph_style(document, name), **tokens)


def set_section_geometry(section, *, landscape: bool = False) -> None:
    # A4, like the source PDF. The side margins leave the 9,360-twip tables room.
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Twips(16838)
        section.page_height = Twips(11906)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Twips(11906)
        section.page_height = Twips(16838)
    section.top_margin = Inches(1)
    section.right_margin = Inches(0.85)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(0.85)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(qn("xml:space"), "preserve")
    instruction_element.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction_element, separate, text, end])


def configure_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(
        "MedTech Europe Code of Ethical Business Practice · September 2024"
    )
    set_run_font(run, size=8.5, color=MUTED)

    section.footer.is_linked_to_previous = False
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(paragraph, "PAGE", "1")


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


# List geometry (twips). Item text starts at LIST_TEXT_INDENT, one step per level.
LIST_TEXT_INDENT = 720
LIST_LEVEL_INDENT = 360
BULLET_HANG = 270
LABEL_HANG = 360
BULLET_TEXT = {"bullet": "•", "square": "■", "dash": "-"}


class NumberingManager:
    """Word bullet lists that print the PDF's own glyph at every level."""

    def __init__(self, document: Document):
        self.numbering = document.part.numbering_part.element
        self.num_ids: dict[str, int] = {}
        self.next_abstract_id = self._next_id("w:abstractNum", "w:abstractNumId")
        self.next_num_id = self._next_id("w:num", "w:numId")

    def _next_id(self, element_name: str, attribute_name: str) -> int:
        values = [
            int(element.get(qn(attribute_name)))
            for element in self.numbering.findall(qn(element_name))
            if element.get(qn(attribute_name)) is not None
        ]
        return max(values, default=0) + 1

    def _abstract(self, glyph: str) -> int:
        abstract_id = self.next_abstract_id
        self.next_abstract_id += 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level in range(3):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), "bullet")
            marker = OxmlElement("w:lvlText")
            marker.set(qn("w:val"), glyph)
            justification = OxmlElement("w:lvlJc")
            justification.set(qn("w:val"), "left")
            paragraph_properties = OxmlElement("w:pPr")
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            position = LIST_TEXT_INDENT + level * LIST_LEVEL_INDENT
            tab.set(qn("w:pos"), str(position))
            tabs.append(tab)
            indentation = OxmlElement("w:ind")
            indentation.set(qn("w:left"), str(position))
            indentation.set(qn("w:hanging"), str(BULLET_HANG))
            paragraph_properties.extend([tabs, indentation])
            lvl.extend([start, num_fmt, marker, justification, paragraph_properties])
            abstract.append(lvl)
        # Word requires every abstractNum before the first num.
        first_num = self.numbering.find(qn("w:num"))
        if first_num is None:
            self.numbering.append(abstract)
        else:
            first_num.addprevious(abstract)
        return abstract_id

    def bullet_list(self, kind: str) -> int:
        if kind not in self.num_ids:
            num_id = self.next_num_id
            self.next_num_id += 1
            num = OxmlElement("w:num")
            num.set(qn("w:numId"), str(num_id))
            abstract_id = OxmlElement("w:abstractNumId")
            abstract_id.set(qn("w:val"), str(self._abstract(BULLET_TEXT[kind])))
            num.append(abstract_id)
            self.numbering.append(num)
            self.num_ids[kind] = num_id
        return self.num_ids[kind]

    @staticmethod
    def apply(paragraph, num_id: int, level: int = 0) -> None:
        paragraph_properties = paragraph._p.get_or_add_pPr()
        num_properties = paragraph_properties.find(qn("w:numPr"))
        if num_properties is None:
            num_properties = OxmlElement("w:numPr")
            paragraph_properties.append(num_properties)
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_properties.extend([ilvl, num])


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=80, start=120, bottom=80, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    properties.append(node)


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(sum(widths_dxa)))
    width.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(indent_dxa))
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)

    for row in table.rows:
        for cell, value in zip(row.cells, widths_dxa):
            cell.width = Inches(value / 1440)
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:w"), str(value))
            cell_width.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), TABLE_BORDER)


def clear_cell(cell) -> None:
    for paragraph in cell.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    cell.paragraphs[0].clear()


def fill_cell(
    cell,
    text: str,
    numbering: NumberingManager,
    *,
    header: bool = False,
) -> None:
    clear_cell(cell)
    lines = [normalize_space(line) for line in (text or "").splitlines() if line.strip()]
    if not lines:
        return
    # "• " and "■ " open bullet items; "◦ " marks a nested bullet, which the PDF
    # also prints as "•".
    markers = {"• ": ("bullet", 0), "◦ ": ("bullet", 1), "■ ": ("square", 0)}
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.style = "Table Header" if header else "Table Text"
        marker = markers.get(plain(line)[:2])
        if marker is not None:
            kind, level = marker
            NumberingManager.apply(paragraph, numbering.bullet_list(kind), level)
            line = strip_visible(line, 2)
        add_marked_runs(paragraph, line)


def add_table(
    document: Document,
    rows: Sequence[Sequence[str]],
    widths_dxa: Sequence[int],
    numbering: NumberingManager,
    *,
    header_rows: int = 1,
    fill: str = LIGHT_BLUE,
) -> object:
    table = document.add_table(rows=len(rows), cols=len(widths_dxa))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            fill_cell(
                table.cell(row_index, column_index),
                value or "",
                numbering,
                header=row_index < header_rows,
            )
            if row_index < header_rows:
                shade_cell(table.cell(row_index, column_index), fill)
        if row_index < header_rows:
            set_repeat_table_header(table.rows[row_index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_note_box(document: Document, title: str, body: str) -> None:
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table)
    cell = table.cell(0, 0)
    shade_cell(cell, NOTE_FILL)
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title)
    set_run_font(run, size=10, color=HEADING_DARK_BLUE, bold=True)
    paragraph = cell.add_paragraph(style="Editorial Note")
    paragraph.add_run(body)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(document: Document) -> object:
    section = document.sections[0]
    section.different_first_page_header_footer = True
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(118)
    paragraph.paragraph_format.space_after = Pt(18)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("MEDTECH EUROPE")
    set_run_font(run, size=11, color=TEAL, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Code of Ethical\nBusiness Practice")
    set_run_font(run, size=30, color=HEADING_DARK_BLUE, bold=True)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(66)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("September 2024")
    set_run_font(run, size=15, color=MUTED)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Editable PDF-derived working copy")
    set_run_font(run, size=10.5, color=MUTED, italic=True)
    return paragraph


def add_front_matter(document: Document) -> None:
    document.add_page_break()
    heading = document.add_heading("Document status and editing guidance", level=1)
    add_note_box(
        document,
        "Provenance",
        (
            f"This working copy was generated only from {SOURCE_PDF.name} "
            f"(SHA-256 {SOURCE_SHA256}). Repeated PDF page furniture and the PDF’s "
            "fixed table-of-contents pagination were converted to Word-native "
            "formatting. No JSON content was used."
        ),
    )
    add_note_box(
        document,
        "Text policy",
        (
            "The source wording is preserved except for the approved corrections "
            "listed in the final Editorial correction record. Clause, list and "
            "footnote labels are typed text, exactly as printed; only bullets use Word "
            "list formatting. That final record and this guidance page are editorial "
            "metadata, not Code text."
        ),
    )
    add_note_box(
        document,
        "Team editing",
        (
            "Edit the Code text directly and use Track Changes. Keep chapter, section, "
            "clause, list-letter, Q-number, A-number, and annex labels intact unless "
            "renumbering is an intended amendment; Word does not renumber them. "
            "PDF-page bookmarks named pdf_page_001 through pdf_page_080 provide stable "
            "source provenance for later JSON updates."
        ),
    )
    heading.paragraph_format.keep_with_next = True

    document.add_page_break()
    document.add_heading("Table of Contents", level=1)
    toc = document.add_paragraph()
    add_field(
        toc,
        'TOC \\o "1-3" \\h \\z \\u',
        "Right-click here and choose Update Field to build the table of contents.",
    )
    document.add_page_break()


MAIN_STYLES = {
    "body": "Code Body",
    "clause": "Code Clause",
    "question": "Code Question",
    "answer": "Code Answer",
    "footnote": "Code Footnote",
    **{kind: "Code List" for kind in LIST_KINDS},
}
# Everything printed in the Q&A column belongs to a question or its answer.
SIDE_STYLES = {
    **MAIN_STYLES,
    "body": "Code Answer",
    "clause": "Code Answer",
    **{kind: "Code Answer List" for kind in LIST_KINDS},
}


def add_marked_runs(paragraph, text: str) -> None:
    bold = superscript = False
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            run = paragraph.add_run("".join(buffer))
            if bold:
                run.bold = True
            if superscript:
                run.font.superscript = True
            buffer.clear()

    for character in text:
        if MARKUP_RE.match(character):
            flush()
            if character == SUP_START:
                superscript = True
            elif character == SUP_END:
                superscript = False
            elif character == BOLD_START:
                bold = True
            else:
                bold = False
            continue
        buffer.append(character)
    flush()


def indent_paragraph(paragraph, level: int, hanging: int = 0) -> None:
    position = LIST_TEXT_INDENT + level * LIST_LEVEL_INDENT
    paragraph_format = paragraph.paragraph_format
    paragraph_format.left_indent = Twips(position)
    if hanging:
        paragraph_format.first_line_indent = Twips(-hanging)
        paragraph_format.tab_stops.add_tab_stop(Twips(position))


def add_block(document: Document, block: Block, numbering: NumberingManager):
    if block.kind == "heading":
        return document.add_heading(block.text, level=block.level)
    styles = SIDE_STYLES if block.column == "side" else MAIN_STYLES
    paragraph = document.add_paragraph(style=styles[block.kind])
    if block.kind == "label":
        add_marked_runs(paragraph, f"{block.label}\t{block.text}")
        indent_paragraph(paragraph, block.level, LABEL_HANG)
    else:
        add_marked_runs(paragraph, block.text)
    if block.kind in BULLET_KINDS:
        NumberingManager.apply(paragraph, numbering.bullet_list(block.kind), block.level)
    elif block.indent is not None:
        indent_paragraph(paragraph, block.indent)
    if block.kind in {"question", "clause"}:
        paragraph.paragraph_format.keep_with_next = True
    return paragraph


def smart_join_split(left: str | None, right: str | None) -> str:
    left = (left or "").rstrip()
    right = (right or "").lstrip()
    if not left:
        return right
    if not right:
        return left
    if left[-1].isalpha() and right[0].islower():
        return left + right
    return f"{left} {right}"


def annex_i_rows(plumber_document) -> list[list[str]]:
    tables = plumber_document.pages[71].find_tables()
    source = max((table.extract() for table in tables), key=len)
    header = [
        (
            "WHICH TYPE OF SUPPORT CAN MEMBER COMPANIES PROVIDE TO WHICH THIRD "
            "PARTY ORGANISED EDUCATIONAL EVENTS?"
        ),
        "",
        (
            "NATIONAL\nThird Party Organised Educational Events attended by "
            "delegates which are local HCPs only)"
        ),
        (
            "INTERNATIONAL\n(Third Party Organised Educational Events attended "
            "by delegates coming from at least two countries of the MedTech "
            f"Europe Geographic Area{SUP_START}1,2{SUP_END})"
        ),
        (
            "INTERNATIONAL\n(Third Party Organised Educational Events attended by "
            "delegates who are Healthcare Professionals registered and practising "
            f"in the MedTech Europe Geographic Area{SUP_START}3{SUP_END})"
        ),
        (
            "INTERNATIONAL\n(Third Party Organised Educational Events to which no "
            "Healthcare Professionals registered and practicing in the MedTech "
            "Europe Geographic Area attend, neither as speakers or delegates)"
        ),
    ]
    rows = [header]
    for row in source[3:]:
        processed = [
            row[0] or "",
            row[1] or "",
            row[2] or "",
            row[3] or "",
            smart_join_split(row[4], row[5]),
            row[6] or "",
        ]
        if "Attendance of Member" in processed[1]:
            # pdfplumber splits and interleaves two words across this cell's
            # artificial extraction columns. Reassemble the printed sentence.
            processed[4] = (
                "Member Company prior internal review against Code principles "
                "required"
            )
        rows.append(processed)
    # Table extraction flattens the two superscript footnote references in the body.
    for pattern, replacement in (
        (r"(EDUCATIONAL GRANTS)4\b", f"\\1{SUP_START}4{SUP_END}"),
        (r"(application of the Code)5\b", f"\\1{SUP_START}5{SUP_END}"),
    ):
        count = 0
        for row in rows[1:]:
            for index, value in enumerate(row):
                row[index], found = re.subn(pattern, replacement, value)
                count += found
        if count != 1:
            raise ValueError(f"Annex I footnote reference {pattern!r} matched {count} times")
    return rows


def annex_ii_rows(plumber_document) -> list[list[str]]:
    page_73 = max(
        (table.extract() for table in plumber_document.pages[72].find_tables()),
        key=len,
    )
    page_74 = max(
        (table.extract() for table in plumber_document.pages[73].find_tables()),
        key=len,
    )
    rows = [
        [
            "IN-KIND CATEGORY",
            "EXAMPLES OF VALUES THAT CAN BE CONSIDERED, VAT EXCLUSIVE WHEN RELEVANT",
        ]
    ]
    goods_value = "\n".join(
        [
            "• Provision of used goods",
            "◦ Fair Market Value",
            "◦ Company book value",
            "• Provision of new goods",
            "◦ For third party goods, the listing price.",
            "◦ Contractual value or Fair Market Value",
            (
                "◦ Internal costs, whether relating to cost of manufacture "
                "or transfer price"
            ),
            "• Loaned goods",
            "◦ Rental equivalent based on depreciation",
            (
                "◦ Rental equivalent to highest-volume rate – NOTE: Rent "
                "cannot exceed accepted values if the equipment were to "
                "be donated or sold"
            ),
            (
                "◦ Exception: if single used products are bundled, they "
                "should be separated and valued as a donation of equipment"
            ),
        ]
    )
    for row in page_73[1:]:
        category = row[0] or ""
        value = smart_join_split(row[1], row[2])
        if category.startswith("Services provided"):
            # The narrow source columns interleave the parenthetical phrases.
            # Preserve the reading order visible in the PDF.
            value = (
                "The salary (inclusive social contributions), or a portion of "
                "the salary, for technical support provided by personnel employed "
                "by the Member Company (based on time spent and salary)"
            )
        elif category.startswith("Goods"):
            value = goods_value
        rows.append([category, value])
    for row in page_74:
        category = row[0] or ""
        value = smart_join_split(row[1], row[2])
        if category.startswith("Goods"):
            # Repair extraction-order artifacts while retaining the PDF's
            # wording and its editable bullet structure.
            value = goods_value
        rows.append([category, value])
    return rows


def annex_vi_rows(plumber_document) -> list[list[str]]:
    source = max(
        (table.extract() for table in plumber_document.pages[77].find_tables()),
        key=len,
    )
    rows = [["Event", "Setting", "Faculty /Speaker", "Delegates"]]
    for row in source[2:]:
        faculty = smart_join_split(row[3], row[4])
        if faculty == "Notallowed":
            faculty = "Not allowed"
        rows.append(
            [
                row[0] or "",
                "\n".join(value for value in (row[1], row[2]) if value),
                faculty,
                row[5] or "",
            ]
        )
    return rows


def add_column_blocks(
    document: Document,
    page_number: int,
    lines: Sequence[Line],
    headings: set[str],
    numbering: NumberingManager,
    corrections: CorrectionLog,
):
    blocks = parse_column(page_number, "main", lines, headings)
    for block in blocks:
        block.text = corrections.apply(page_number, block.text)
    promote_subheadings(blocks)
    assign_levels(blocks, {(page_number, "main"): column_text_left(lines)})
    first = None
    for block in blocks:
        paragraph = add_block(document, block, numbering)
        first = first or paragraph
    return first


def corrected_rows(
    rows: list[list[str]], pages: Sequence[int], corrections: CorrectionLog
) -> list[list[str]]:
    for row in rows:
        for index, value in enumerate(row):
            for page_number in pages:
                value = corrections.apply(page_number, value)
            row[index] = value
    return rows


def add_annex_i(
    document: Document,
    lines: Sequence[Line],
    plumber_document,
    numbering: NumberingManager,
    corrections: CorrectionLog,
    bookmark_id: int,
) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(section, landscape=True)
    configure_header_footer(section)
    heading = document.add_heading("ANNEX I", level=2)
    add_bookmark(heading, "pdf_page_072", bookmark_id)
    document.add_heading("CVS scope: When are CVS assessments required?", level=3)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("PRIOR CVS SUBMISSION")
    set_run_font(run, size=10, color=HEADING_DARK_BLUE, bold=True)
    group_table = document.add_table(rows=1, cols=2)
    set_table_geometry(group_table, [6480, 6480], indent_dxa=0)
    fill_cell(
        group_table.cell(0, 0),
        "IN MEDTECH EUROPE GEOGRAPHIC AREA",
        numbering,
        header=True,
    )
    fill_cell(
        group_table.cell(0, 1),
        "OUTSIDE MEDTECH EUROPE GEOGRAPHIC AREA",
        numbering,
        header=True,
    )
    add_table(
        document,
        corrected_rows(annex_i_rows(plumber_document), [72], corrections),
        [1500, 2100, 2250, 2250, 2400, 2460],
        numbering,
    )
    footnote_start = next(
        (index for index, line in enumerate(lines) if plain(line.text).startswith("1) ")),
        len(lines),
    )
    add_column_blocks(
        document, 72, lines[footnote_start:], set(), numbering, corrections
    )

    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(section)
    configure_header_footer(section)


def add_annex_ii(
    document: Document,
    page_73_lines: Sequence[Line],
    page_74_lines: Sequence[Line],
    headings: set[str],
    plumber_document,
    numbering: NumberingManager,
    corrections: CorrectionLog,
    bookmark_id_73: int,
    bookmark_id_74: int,
) -> None:
    heading = document.add_heading(
        "ANNEX II Calculating the value of In Kind Educational Grants", level=2
    )
    add_bookmark(heading, "pdf_page_073", bookmark_id_73)
    table_start = next(
        (
            index
            for index, line in enumerate(page_73_lines)
            if plain(line.text).startswith("IN-KIND CATEGORY")
        ),
        len(page_73_lines),
    )
    intro = [line for line in page_73_lines[:table_start] if line.font != "title"]
    add_column_blocks(document, 73, intro, headings, numbering, corrections)
    add_table(
        document,
        corrected_rows(annex_ii_rows(plumber_document), [73, 74], corrections),
        [2700, 6660],
        numbering,
    )
    marker = document.add_paragraph(style="Code Body")
    add_bookmark(marker, "pdf_page_074", bookmark_id_74)
    marker.add_run("In-Kind Educational Grant do not include the following:")
    postamble_start = next(
        (
            index
            for index, line in enumerate(page_74_lines)
            if plain(line.text).startswith("In-Kind Educational Grant do not include")
        ),
        len(page_74_lines),
    )
    add_column_blocks(
        document,
        74,
        page_74_lines[postamble_start + 1 :],
        headings,
        numbering,
        corrections,
    )


def add_annex_iii(
    document: Document,
    lines: Sequence[Line],
    numbering: NumberingManager,
    bookmark_id: int,
) -> None:
    heading = document.add_heading("Annex III", level=2)
    add_bookmark(heading, "pdf_page_075", bookmark_id)
    document.add_heading(
        "The Geographical Area where the Code applies as the minimum standard",
        level=3,
    )
    document.add_paragraph(
        "The MedTech Europe Geographic Area currently includes",
        style="Code Body",
    )
    document.add_heading("Countries with National Associations:", level=3)
    country_rows = []
    for line in lines:
        text = plain(line.text)
        if not text.startswith("■ "):
            continue
        values = [f"■ {normalize_space(value)}" for value in text.split("■") if value.strip()]
        country_rows.append((values + [""])[:2])
    split_index = next(
        (
            index
            for index, row in enumerate(country_rows)
            if row[0] == "■ Iceland"
        ),
        len(country_rows),
    )
    add_table(
        document,
        country_rows[:split_index],
        [4680, 4680],
        numbering,
        header_rows=0,
        fill=LIGHT_GRAY,
    )
    document.add_heading(
        (
            "Countries party to the European Economic Area agreement without a "
            "MedTech Europe National Association:"
        ),
        level=3,
    )
    for value in [item for row in country_rows[split_index:] for item in row if item]:
        add_block(document, Block("square", value[2:], 75), numbering)
    footnote = next(
        (plain(line.text) for line in lines if plain(line.text).startswith("*Countries")),
        "",
    )
    if footnote:
        document.add_paragraph(footnote, style="Code Footnote")


def add_annex_vi(
    document: Document,
    lines: Sequence[Line],
    headings: set[str],
    plumber_document,
    numbering: NumberingManager,
    corrections: CorrectionLog,
    bookmark_id: int,
) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(section, landscape=True)
    configure_header_footer(section)
    heading = document.add_heading(
        "ANNEX VI Direct support to HCP participation in Events", level=2
    )
    add_bookmark(heading, "pdf_page_078", bookmark_id)
    document.add_heading("Direct Support for HCP attendance", level=3)
    add_table(
        document,
        corrected_rows(annex_vi_rows(plumber_document), [78], corrections),
        [2200, 5000, 3000, 2760],
        numbering,
    )
    description_start = next(
        (index for index, line in enumerate(lines) if plain(line.text) == "Description:"),
        len(lines),
    )
    add_column_blocks(
        document, 78, lines[description_start:], headings, numbering, corrections
    )

    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_section_geometry(section)
    configure_header_footer(section)


def add_correction_record(
    document: Document, numbering: NumberingManager
) -> None:
    document.add_page_break()
    document.add_heading(
        "Editorial correction record (not part of the Code)", level=1
    )
    document.add_paragraph(
        (
            "Only the page-scoped corrections below differ from the source PDF’s "
            "visible wording. They fix typographical, grammatical and punctuation "
            "defects and were approved as not changing the intended normative meaning."
        ),
        style="Editorial Note",
    )
    rows = [["PDF page", "Source text", "Corrected text", "Reason"]]
    rows.extend(
        [
            str(item.page),
            item.original,
            item.replacement.replace("\n", " "),
            item.reason,
        ]
        for item in CORRECTIONS
    )
    add_table(
        document,
        rows,
        [900, 3000, 3360, 2100],
        numbering,
        fill=LIGHT_GRAY,
    )


def structural_audit(document: Document) -> None:
    required_styles = {
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Code Body",
        "Code Question",
        "Code Answer",
        "Code List",
        "Code Footnote",
    }
    missing = required_styles - {style.name for style in document.styles}
    if missing:
        raise ValueError(f"Missing required styles: {sorted(missing)}")
    if len(document.tables) < 7:
        raise ValueError(
            f"Expected editable annex/metadata tables; found {len(document.tables)}"
        )
    correction_heading = "Editorial correction record (not part of the Code)"
    body_paragraphs = []
    for paragraph in document.paragraphs:
        if paragraph.text == correction_heading:
            break
        if paragraph.text:
            body_paragraphs.append(paragraph.text)
    body_tables = document.tables[:-1] if document.tables else []
    text = "\n".join(
        [
            *body_paragraphs,
            *(
                paragraph.text
                for table in body_tables
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
                if paragraph.text
            ),
        ]
    )
    forbidden = (
        "July 2026 Update",
        "Knowledge Quiz",
        "TPPT Checker",
        "Summary:",
    )
    found = [value for value in forbidden if value in text]
    if found:
        raise ValueError(f"Non-PDF/app-derived content leaked into DOCX: {found}")
    for correction in CORRECTIONS:
        if re.search(correction.pattern, text, correction.flags):
            raise ValueError(
                f"Uncorrected source typo remains in DOCX: {correction.original}"
            )
    for page_number in range(1, 81):
        bookmark = f"pdf_page_{page_number:03d}"
        if bookmark.encode() not in document.part.blob:
            raise ValueError(f"Missing source-page bookmark: {bookmark}")
    if MARKUP_RE.search(text):
        raise ValueError("Internal markup characters leaked into the DOCX text.")
    body_runs = [
        run
        for paragraph in [
            *document.paragraphs,
            *(
                paragraph
                for table in body_tables
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
            ),
        ]
        for run in paragraph.runs
    ]
    superscripts = [run.text for run in body_runs if run.font.superscript]
    if len(superscripts) != EXPECTED_SUPERSCRIPTS:
        raise ValueError(
            f"Expected {EXPECTED_SUPERSCRIPTS} superscript footnote references; "
            f"found {superscripts}"
        )


# Pages built from parsed text; the other annex pages have dedicated table handlers.
TEXT_PAGES = (2, *range(5, 72), 76, 77, 79, 80)
EXPECTED_SUPERSCRIPTS = 9


def build(source_pdf: Path, output_docx: Path) -> None:
    if sha256(source_pdf) != SOURCE_SHA256:
        raise ValueError(
            "Source PDF hash does not match the reviewed September 2024 source."
        )
    corrections = CorrectionLog()

    document = Document()
    configure_styles(document)
    set_section_geometry(document.sections[0])
    configure_header_footer(document.sections[0])
    document.core_properties.title = (
        "MedTech Europe Code of Ethical Business Practice — September 2024"
    )
    document.core_properties.subject = "Editable PDF-derived working copy"
    document.core_properties.author = "MedTech Europe"
    document.core_properties.comments = (
        f"Generated solely from {source_pdf.name}; SHA-256 {SOURCE_SHA256}."
    )
    numbering = NumberingManager(document)
    bookmark_ids = iter(range(1, 1000))

    cover_anchor = add_cover(document)
    add_bookmark(cover_anchor, "pdf_page_001", next(bookmark_ids))
    add_front_matter(document)
    toc_heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "Table of Contents"
    )
    add_bookmark(toc_heading, "pdf_page_003", next(bookmark_ids))
    add_bookmark(toc_heading, "pdf_page_004", next(bookmark_ids))

    with pdfplumber.open(source_pdf) as plumber_document:
        if len(plumber_document.pages) != 80:
            raise ValueError(f"Expected 80 PDF pages; found {len(plumber_document.pages)}")
        headings = extract_toc_headings(plumber_document)
        columns = {
            page_number: page_columns(plumber_document.pages[page_number - 1], page_number)
            for page_number in range(2, 81)
            if page_number not in (3, 4)
        }

        pages: dict[int, dict[str, list[Block]]] = {}
        text_left: dict[tuple[int, str], float] = {}
        for page_number in TEXT_PAGES:
            pages[page_number] = {}
            for column, lines in columns[page_number].items():
                if page_number in MAJOR_TITLES:
                    if column != "main":
                        continue
                    lines = strip_cover_lines(page_number, lines)
                text_left[(page_number, column)] = column_text_left(lines)
                blocks = parse_column(page_number, column, lines, headings)
                for block in blocks:
                    block.text = corrections.apply(page_number, block.text)
                pages[page_number][column] = blocks
        # Text never continues onto a chapter title page.
        merge_page_breaks(pages, fresh_starts=set(MAJOR_TITLES))
        promote_subheadings(
            block for page in pages.values() for blocks in page.values() for block in blocks
        )
        for column in ("main", "side", "note"):
            assign_levels(
                [
                    block
                    for page_number in TEXT_PAGES
                    for block in pages[page_number].get(column, [])
                ],
                text_left,
            )
        own_blocks = {
            page_number: any(pages[page_number].get(column) for column in ("main", "side", "note"))
            or page_number in MAJOR_TITLES
            for page_number in TEXT_PAGES
        }

        for page_number in range(2, 81):
            if page_number in (3, 4, 74):
                continue
            if page_number == 72:
                add_annex_i(
                    document,
                    columns[72]["main"],
                    plumber_document,
                    numbering,
                    corrections,
                    next(bookmark_ids),
                )
                continue
            if page_number == 73:
                add_annex_ii(
                    document,
                    columns[73]["main"],
                    columns[74]["main"],
                    headings,
                    plumber_document,
                    numbering,
                    corrections,
                    next(bookmark_ids),
                    next(bookmark_ids),
                )
                continue
            if page_number == 75:
                add_annex_iii(document, columns[75]["main"], numbering, next(bookmark_ids))
                continue
            if page_number == 78:
                add_annex_vi(
                    document,
                    columns[78]["main"],
                    headings,
                    plumber_document,
                    numbering,
                    corrections,
                    next(bookmark_ids),
                )
                continue

            anchor = None
            if page_number in MAJOR_TITLES:
                document.add_page_break()
                anchor = document.add_heading(MAJOR_TITLES[page_number], level=1)
            for column in ("main", "side", "note"):
                for block in pages[page_number].get(column, []):
                    paragraph = add_block(document, block, numbering)
                    anchor = anchor or paragraph
                    for merged_page in block.pages:
                        # A page whose only text continued a paragraph from the
                        # previous page is anchored on that paragraph.
                        if not own_blocks.get(merged_page, True):
                            add_bookmark(
                                paragraph, f"pdf_page_{merged_page:03d}", next(bookmark_ids)
                            )
            if anchor is None:
                if own_blocks[page_number]:
                    raise ValueError(f"Page {page_number}: no content blocks")
                continue
            add_bookmark(anchor, f"pdf_page_{page_number:03d}", next(bookmark_ids))

    add_correction_record(document, numbering)
    corrections.assert_complete()
    structural_audit(document)
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)

    reopened = Document(output_docx)
    structural_audit(reopened)
    print(
        f"Created {output_docx} from {source_pdf.name}: "
        f"{len(reopened.paragraphs)} paragraphs, {len(reopened.tables)} tables, "
        f"{len(reopened.sections)} sections, {len(CORRECTIONS)} corrections."
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE_PDF)
    parser.add_argument("--output", type=Path, default=OUTPUT_DOCX)
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    build(arguments.source.resolve(), arguments.output.resolve())
